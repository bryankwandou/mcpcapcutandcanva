"""Original template generators: your own media + text in, a finished design out.

* CapCut templates build a complete timeline (clips, captions, title, CTA,
  music) as a draft you can fine-tune in the live editor.
* Canva templates are generated as PPTX with real, separate elements (text
  boxes, shapes, photos). Importing that PPTX into Canva gives a design whose
  every element is editable and movable.
"""
from __future__ import annotations

import time
from pathlib import Path
from typing import Any

from . import capcut, scene

PALETTES = {
    "bold": {"bg": "#111111", "primary": "#FFD400", "text": "#FFFFFF", "accent": "#FF3B30"},
    "fresh": {"bg": "#E8F7F0", "primary": "#0E9F6E", "text": "#0B2E22", "accent": "#FF8A3D"},
    "elegant": {"bg": "#F6F1EA", "primary": "#8C6A43", "text": "#2B2420", "accent": "#C9A66B"},
    "neon": {"bg": "#0B0B1E", "primary": "#00E5FF", "text": "#FFFFFF", "accent": "#FF2BD6"},
    "pastel": {"bg": "#FFF4F7", "primary": "#F27BA6", "text": "#3A2A33", "accent": "#7CC4F2"},
}

CAPCUT_TEMPLATES = {
    "promo": "Promo/iklan 9:16: judul pembuka, klip dengan caption, penutup CTA, musik.",
    "slideshow": "Slideshow foto/video dengan caption per slide dan musik.",
    "quotes": "Quotes/lirik: satu background, beberapa baris teks bergantian.",
    "youtube_intro": "Intro YouTube 16:9: judul besar + subjudul di atas klip.",
}
CANVA_TEMPLATES = {
    "instagram_post": (1080, 1350, "Poster feed Instagram 4:5"),
    "story": (1080, 1920, "Story/Reels cover 9:16"),
    "youtube_thumbnail": (1280, 720, "Thumbnail YouTube 16:9"),
    "presentation": (1920, 1080, "Presentasi: cover + slide isi + penutup"),
}


def list_templates() -> dict:
    return {"capcut": CAPCUT_TEMPLATES,
            "canva": {k: v[2] for k, v in CANVA_TEMPLATES.items()},
            "palettes": list(PALETTES)}


def _kind(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext in {".mp3", ".wav", ".m4a", ".aac", ".ogg"}:
        return "audio"
    return "photo" if ext in {".jpg", ".jpeg", ".png", ".webp", ".gif"} else "video"


# ------------------------------------------------------------- CapCut ----
def build_capcut(template: str, name: str, media: list[str], title: str = "",
                 captions: list[str] | None = None, cta: str = "", music: str | None = None,
                 palette: str = "bold", clip_seconds: float = 2.5,
                 base_draft: str | None = None, style: dict[str, str] | None = None) -> dict:
    """style (optional, keys from capcut_library_list): {"transition": key, "filter": key,
    "font": key, "text_animation": key, "clip_animation": key, "sticker": key, "effect": key}"""
    if template not in CAPCUT_TEMPLATES:
        raise ValueError(f"Template tidak dikenal. Pilihan: {list(CAPCUT_TEMPLATES)}")
    pal = PALETTES.get(palette, PALETTES["bold"])
    w, h = (1920, 1080) if template == "youtube_intro" else (1080, 1920)
    capcut.create_draft(name, w, h, template=base_draft)
    captions = captions or []
    vid, txt, aud = capcut._uid(), capcut._uid(), capcut._uid()
    els: list[dict[str, Any]] = []

    def clip(src, start, dur, **kw):
        els.append({"id": capcut._uid(), "track": vid, "type": _kind(src), "src": str(Path(src).expanduser().resolve()),
                    "start": start, "duration": dur, "x": 0.0, "y": 0.0, "scale": 1.0, "rotation": 0.0,
                    "alpha": 1.0, "volume": 1.0, **kw})

    def text(t, start, dur, y=0.0, size=8.0, color=None):
        els.append({"id": capcut._uid(), "track": txt, "type": "text", "text": t, "start": start,
                    "duration": dur, "x": 0.0, "y": y, "scale": 1.0, "rotation": 0.0, "alpha": 1.0,
                    "volume": 1.0, "color": color or pal["text"], "font_size": size})

    t = 0.0
    if template == "promo":
        offset = 1 if title else 0  # first clip carries the title, the rest carry captions
        for i, m in enumerate(media):
            clip(m, t, clip_seconds)
            if i == 0 and title:
                text(title, t, clip_seconds, y=0.55, size=14, color=pal["primary"])
            elif i - offset < len(captions):
                text(captions[i - offset], t, clip_seconds, y=-0.6, size=9)
            t += clip_seconds
        if cta:
            last = media[-1] if media else None
            if last:
                clip(last, t, 2.0, scale=1.1, alpha=0.6)
            text(cta, t, 2.0, y=0.0, size=13, color=pal["accent"])
            t += 2.0
    elif template == "slideshow":
        if title:
            text(title, 0, clip_seconds, y=0.6, size=12, color=pal["primary"])
        for i, m in enumerate(media):
            clip(m, t, clip_seconds)
            if i < len(captions):
                text(captions[i], t, clip_seconds, y=-0.65, size=8)
            t += clip_seconds
    elif template == "quotes":
        lines = captions or [title]
        per = max(clip_seconds, 3.0)
        total = per * len(lines)
        if media:
            clip(media[0], 0, total, alpha=0.75)
        for i, line in enumerate(lines):
            text(line, i * per, per, y=0.0, size=10)
        if title and captions:
            text(title, 0, total, y=-0.8, size=6, color=pal["primary"])
        t = total
    elif template == "youtube_intro":
        dur = max(clip_seconds * max(1, len(media)), 4.0)
        for i, m in enumerate(media):
            clip(m, i * dur / len(media), dur / len(media))
        text(title or "JUDUL CHANNEL", 0.3, dur - 0.3, y=0.1, size=16, color=pal["primary"])
        if captions:
            text(captions[0], 0.8, dur - 0.8, y=-0.25, size=8)
        t = dur
    tracks = [{"id": vid, "type": "video"}, {"id": txt, "type": "text"}]
    if music:
        tracks.append({"id": aud, "type": "audio"})
        els.append({"id": capcut._uid(), "track": aud, "type": "audio",
                    "src": str(Path(music).expanduser().resolve()), "start": 0, "duration": t,
                    "volume": 0.8, "x": 0, "y": 0, "scale": 1, "rotation": 0, "alpha": 1})
    res = scene.apply_to_draft({"draft": name, "tracks": tracks, "elements": els})
    if style:
        res["style"] = _apply_style(name, style, els, t)
    return {"draft": name, "duration": t, **res}


def _apply_style(name: str, style: dict[str, str], els: list[dict], total: float) -> list:
    """Decorate a generated project with elements from your personal CapCut library."""
    from . import capcut_library as lib

    done = []
    clips = sorted((e for e in els if e["type"] in ("video", "photo")), key=lambda e: e["start"])
    texts = [e for e in els if e["type"] == "text"]
    jobs = []
    if style.get("transition"):
        jobs += [(style["transition"], {"segment_id": c["id"]}) for c in clips[:-1]]
    if style.get("clip_animation"):
        jobs += [(style["clip_animation"], {"segment_id": c["id"]}) for c in clips]
    if style.get("font"):
        jobs += [(style["font"], {"segment_id": e["id"]}) for e in texts]
    if style.get("text_animation"):
        jobs += [(style["text_animation"], {"segment_id": e["id"]}) for e in texts]
    if style.get("filter"):
        jobs.append((style["filter"], {"start": 0, "duration": total}))
    if style.get("effect"):
        jobs.append((style["effect"], {"start": 0, "duration": min(2.0, total)}))
    if style.get("sticker"):
        jobs.append((style["sticker"], {"start": 0, "duration": total, "x": 0.6, "y": 0.75, "scale": 0.5}))
    for key, kw in jobs:
        try:
            done.append(lib.apply(name, key, **kw))
        except Exception as e:  # keep building even if one element can't be applied
            done.append({"ok": False, "key": key, "error": str(e)})
    return done


# -------------------------------------------------------------- Canva ----
def _rgb(hexs: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexs.lstrip("#").upper())


def build_pptx(template: str, out_path: str, title: str, subtitle: str = "",
               body: list[str] | None = None, images: list[str] | None = None,
               cta: str = "", palette: str = "bold", font: str = "Montserrat") -> str:
    """Create an editable PPTX design. Import it with canva_import_file."""
    from PIL import Image as PILImage
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu

    if template not in CANVA_TEMPLATES:
        raise ValueError(f"Template tidak dikenal. Pilihan: {list(CANVA_TEMPLATES)}")
    W, H, _ = CANVA_TEMPLATES[template]
    pal = PALETTES.get(palette, PALETTES["bold"])
    px = lambda v: Emu(int(v * 9525))  # noqa: E731  (1px at 96dpi)
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    blank = prs.slide_layouts[6]
    images = [str(Path(i).expanduser()) for i in (images or [])]
    body = body or []

    def slide():
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = _rgb(pal["bg"])
        return s

    def rect(s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
        r = s.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
        r.fill.solid()
        r.fill.fore_color.rgb = _rgb(color)
        r.line.fill.background()
        return r

    def text(s, t, x, y, w, h, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for i, line in enumerate(t.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = line
            r.font.size = Emu(int(size * 9525))
            r.font.bold = bold
            r.font.name = font
            r.font.color.rgb = _rgb(color)
        return tb

    def photo(s, path, x, y, w, h):
        """Place an image cropped to fill (object-fit: cover)."""
        pic = s.shapes.add_picture(path, px(x), px(y), px(w), px(h))
        iw, ih = PILImage.open(path).size
        box, img = w / h, iw / ih
        if img > box:
            c = (1 - box / img) / 2
            pic.crop_left = pic.crop_right = c
        else:
            c = (1 - img / box) / 2
            pic.crop_top = pic.crop_bottom = c
        return pic

    m = round(min(W, H) * 0.06)
    if template in ("instagram_post", "story"):
        s = slide()
        ph = H * (0.58 if template == "instagram_post" else 0.55)
        if images:
            photo(s, images[0], 0, 0, W, ph)
        else:
            rect(s, 0, 0, W, ph, pal["primary"])
        rect(s, m, ph - 40, W * 0.35, 18, pal["accent"])
        text(s, title, m, ph + m * 0.6, W - 2 * m, H * 0.16, W * 0.075, pal["text"], bold=True)
        if subtitle:
            text(s, subtitle, m, ph + m * 0.6 + H * 0.15, W - 2 * m, H * 0.08, W * 0.035, pal["text"])
        if body:
            text(s, "\n".join("• " + b for b in body), m, ph + H * 0.27, W - 2 * m, H * 0.12,
                 W * 0.03, pal["text"])
        if cta:
            b = rect(s, m, H - m - H * 0.065, W * 0.5, H * 0.065, pal["primary"], MSO_SHAPE.ROUNDED_RECTANGLE)
            text(s, cta, m, H - m - H * 0.065, W * 0.5, H * 0.065, W * 0.035, pal["bg"], bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            _ = b
    elif template == "youtube_thumbnail":
        s = slide()
        if images:
            photo(s, images[0], W * 0.45, 0, W * 0.55, H)
        rect(s, 0, 0, W * 0.5, H, pal["bg"])
        rect(s, W * 0.5 - 12, 0, 24, H, pal["primary"])
        text(s, title.upper(), m, m, W * 0.46 - m, H * 0.62, W * 0.075, pal["text"], bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
        if subtitle:
            tag = rect(s, m, H - m - 80, W * 0.32, 80, pal["accent"], MSO_SHAPE.ROUNDED_RECTANGLE)
            text(s, subtitle, m, H - m - 80, W * 0.32, 80, 40, "#FFFFFF", bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            _ = tag
    else:  # presentation
        s = slide()
        if images:
            photo(s, images[0], W * 0.5, 0, W * 0.5, H)
        rect(s, m, H * 0.3, 14, H * 0.4, pal["primary"])
        text(s, title, m + 50, H * 0.28, W * 0.42, H * 0.3, 88, pal["text"], bold=True, anchor=MSO_ANCHOR.BOTTOM)
        if subtitle:
            text(s, subtitle, m + 50, H * 0.6, W * 0.42, H * 0.15, 36, pal["text"])
        for i, point in enumerate(body):
            s = slide()
            rect(s, 0, 0, W, 16, pal["primary"])
            heading, _, detail = point.partition(":")
            text(s, f"{i + 1:02d}", m, m * 1.5, 200, 120, 72, pal["accent"], bold=True)
            text(s, heading.strip(), m, m * 1.5 + 130, W * 0.5 - m, 200, 64, pal["text"], bold=True)
            if detail:
                text(s, detail.strip(), m, m * 1.5 + 340, W * 0.5 - m, H * 0.4, 34, pal["text"])
            img = images[(i + 1) % len(images)] if len(images) > 1 else None
            if img:
                photo(s, img, W * 0.55, m * 1.5, W * 0.45 - m, H - 3 * m)
            else:
                rect(s, W * 0.55, m * 1.5, W * 0.45 - m, H - 3 * m, pal["primary"], MSO_SHAPE.ROUNDED_RECTANGLE)
        s = slide()
        text(s, cta or "Terima kasih!", m, 0, W - 2 * m, H, 96, pal["primary"], bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    out = Path(out_path).expanduser()
    if out.is_dir() or not out.suffix:
        out = out / f"{template}-{int(time.time())}.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return str(out)
